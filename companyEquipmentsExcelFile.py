import requests
import time
from io import BytesIO
from collections import Counter
import openpyxl

# Credenciais da aplicação
CLIENT_ID = "0b9754b4-5b19-48bc-8ea7-3e44322c0bf5"
TENANT_ID = "9f40026c-65e9-49da-a956-531b02c30f6f"
#CLIENT_OITROCARPORSEGREDOEMINGLES = "lCy8Q~LbfNQFMCM6UhMjE6q.yjGpa3EApHdLFchR"

# URL para obter o token
TOKEN_URL = f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token"

# Variáveis globais para armazenar o token e o tempo de expiração
access_token = None
token_expiration_time = 0

def get_access_token():
    """
    Obtém um novo token de acesso para a API Microsoft Graph.
    Verifica se o token atual expirou e, se necessário, obtém um novo.
    """
    global access_token, token_expiration_time

    # Se o token expirou ou ainda não foi obtido, obter um novo
    if access_token is None or time.time() >= token_expiration_time:
        data = {
            "grant_type": "client_credentials",
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
            "scope": "https://graph.microsoft.com/.default"
        }

        try:
            response = requests.post(TOKEN_URL, data=data)
            if response.status_code == 200:
                token_data = response.json()
                access_token = token_data.get("access_token")
                # A duração do token é geralmente de 1 hora (3600 segundos)
                token_expiration_time = time.time() + token_data.get("expires_in", 3600)
                return access_token
            else:
                raise Exception(f"Erro ao obter token: {response.status_code} - {response.json().get('error_description', response.text)}")
        except Exception as e:
            raise Exception(f"Falha ao tentar obter o token de acesso: {str(e)}")
    else:
        return access_token


def process_excel_data():
    """
    Faz download de um arquivo do OneDrive e processa os dados em formato Excel,
    lendo da aba 'CORRETA' e contando os valores na coluna B apenas se o valor da coluna G for 'PATIO'.
    """
    try:
        access_token = get_access_token()

        timestamp = int(time.time())
        url = f'https://graph.microsoft.com/v1.0/drives/b!MksJNZWQn0GCLzJ4pzJnxi_wOlIlFIVPnCiQAI_qFP3U7j9DlkvuS6COUrMPB_5m/items/01JVDEGKMXSHFPAD3RXZE24CIZYDSOULMW/content?timestamp={timestamp}'

        headers = {
            'Authorization': f'Bearer {access_token}',
            'Cache-Control': 'no-cache',
            'Pragma': 'no-cache'
        }

        response = requests.get(url, headers=headers)

        if response.status_code == 200:
            print("Download do arquivo concluído com sucesso.")

            file_content = BytesIO(response.content)
            workbook = openpyxl.load_workbook(file_content, data_only=True)

            if 'CORRETA' not in workbook.sheetnames:
                raise Exception("A aba 'CORRETA' não foi encontrada no arquivo Excel.")

            sheet = workbook['CORRETA']

            values = []

            for row_num, row in enumerate(sheet.iter_rows(min_row=3, max_col=9, values_only=True), start=3):
                value_b = row[1]  # MODELO
                value_g = row[6]  # LOCAL

                print(f"Linha {row_num}: MODELO={value_b}, LOCAL={value_g}")

                if value_b and value_g and str(value_g).strip().lower() == 'patio':
                    value_b = str(value_b).strip()
                    values.append(value_b)

            print(f"Valores filtrados na coluna B: {values}")

            counts = Counter(values)

            result = dict(counts)

            print(f"Resultado final: {result}")

            return result

        else:
            error_msg = f"Erro ao baixar o arquivo: {response.status_code} - {response.text}"
            print(error_msg)
            raise Exception(error_msg)

    except Exception as e:
        error_result = {"error": f"Erro durante o processamento dos dados: {str(e)}"}
        print(error_result)
        return error_result